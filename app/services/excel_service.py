from dataclasses import dataclass
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


STATIC_PATH = Path.cwd() / "static"


@dataclass
class ProjectionRow:
    code: str | None
    productName: str
    brandName: str
    unitName: str
    currentStock: int
    averageDay: float
    projectionDays: int
    averageCost: float
    classification: str

    def to_sheet_row(self, row_index: int):
        return [
            self.code,
            self.productName,
            self.brandName,
            self.unitName,
            self.currentStock,
            self.averageDay,  # F
            self.projectionDays,  # G
            "=" + f"F{row_index}" + "*" + f"G{row_index}" + "-" + f"E{row_index}",  # H
            self.averageCost,  # I
            "=" + f"H{row_index}" + "*" + f"I{row_index}",
            self.classification,
        ]


@dataclass
class ProjectionSheet:
    name: str
    title: str
    subtitle: str
    rows: list[ProjectionRow]


class ProjectionExcelService:
    def __init__(self):
        pass

    def create_projection_xslx(self, sheets: list[ProjectionSheet]) -> Path:
        workbook: Workbook = Workbook()

        for sheet in sheets:
            worksheet: Worksheet = workbook.create_sheet(sheet.name)
            worksheet.title = sheet.title
            worksheet.append(sheet.title.split("\n"))
            worksheet.append(sheet.subtitle.split("\n"))
            worksheet.append(
                [
                    "Código",
                    "Producto",
                    "Marca",
                    "Unidad",
                    "Stock Actual",
                    "Días Promedio",
                    "Días a Proyectar",
                    "Stock Proyectado",
                    "Costo Promedio",
                    "Costo Estimado",
                    "Clasificación",
                ]
            )
            for row in sheet.rows:
                current_row_index = worksheet.max_row
                next_row_index = current_row_index + 1
                worksheet.append(row.to_sheet_row(next_row_index))

        if "Sheet1" in workbook.sheetnames:
            workbook.remove(workbook["Sheet1"])

        if "Sheet" in workbook.sheetnames:
            workbook.remove(workbook["Sheet"])

        workbook.save(STATIC_PATH / "projection.xlsx")

        return STATIC_PATH / "projection.xlsx"

    def get_projection_xslx(self) -> Path:
        return STATIC_PATH / "projection.xlsx"


@dataclass
class PriceListRow:
    code: str | None
    productName: str
    brandName: str
    unitName: str
    salePrice: float
    salePrice2: float
    salePrice3: float

    def to_sheet_row(self):
        return [
            self.code,
            self.productName,
            self.brandName,
            self.unitName,
            self.salePrice,
            self.salePrice2,
            self.salePrice3,
        ]


@dataclass
class PriceListSheet:
    name: str
    title: str
    subtitle: str
    rows: list[PriceListRow]


class PriceListService:
    def __init__(self):
        pass

    def create_price_list_xslx(self, sheets: list[PriceListSheet]) -> Path:
        workbook: Workbook = Workbook()

        for sheet in sheets:
            worksheet: Worksheet = workbook.create_sheet(sheet.name)
            worksheet.title = sheet.title
            worksheet.append(sheet.title.split("\n"))
            worksheet.append(sheet.subtitle.split("\n"))
            worksheet.append(
                [
                    "Código",
                    "Producto",
                    "Marca",
                    "Unidad",
                    "Precio 1",
                    "Precio 2",
                    "Precio 3",
                ]
            )
            for row in sheet.rows:
                worksheet.append(row.to_sheet_row())

        if "Sheet1" in workbook.sheetnames:
            workbook.remove(workbook["Sheet1"])

        if "Sheet" in workbook.sheetnames:
            workbook.remove(workbook["Sheet"])

        workbook.save(STATIC_PATH / "price_list.xlsx")

        return STATIC_PATH / "price_list.xlsx"

    def get_price_list_xslx(self) -> Path:
        return STATIC_PATH / "price_list.xlsx"
